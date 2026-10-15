import csv
from datetime import datetime
from io import BytesIO, StringIO

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def tabular_export(request, filename, headers, rows):
    export_format = request.GET.get("export", "").lower()
    if export_format not in {"csv", "xlsx"}:
        return None

    rows = list(rows)
    if settings.DEMO_MODE:
        filename = f"DEMO-{filename}"
        headers = ["Tipo de datos", *headers]
        rows = [["Demostración: gestiones y estados simulados", *row] for row in rows]
    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows([[excel_safe(value) for value in row] for row in rows])
        response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        return response

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="255B8E")
    for row in rows:
        sheet.append([excel_safe(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 60)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def excel_safe(value):
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
