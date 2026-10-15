from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl.utils import column_index_from_string

from .carga import emails, phones, prepare_company, whatsapp_url
from .management.commands.importar_carga import LAYOUT, Command
from .models import Actividad, Critica, Empresa, EstadoContacto, EstadoEncuesta


def example_row(**overrides):
    return {"E": "Lima", "F": "20000000001", "G": "Empresa de prueba",
            "H": "7492", "N": 0, "Q": 0, "W": 0, **overrides}


class CargaParsingTests(SimpleTestCase):
    def test_prior_years_are_separate_and_current_zero_does_not_block_fallback(self):
        row = example_row(AE="Persona de 2025", AF="Cargo 2025", AH="987654321",
                          AQ="Persona de 2024", AR="Cargo 2024", AT="976543210")
        values, source = prepare_company(row)
        self.assertEqual(values["contacto_principal_nombre"], "Persona de 2025")
        self.assertEqual(values["contacto_principal_cargo"], "Cargo 2025")
        self.assertEqual(values["contacto_principal_telefono"], "987654321")
        self.assertEqual(values["contacto_alternativo_nombre"], "Persona de 2024")
        self.assertIn("AE", source["origenes"]["contacto_principal_nombre"])
        self.assertEqual(source["columnas"]["N"], "0")

    def test_different_peoples_roles_are_not_combined(self):
        values, _ = prepare_company(example_row(N="Persona actual", AE="Otra persona", AF="Otro cargo"))
        self.assertEqual(values["contacto_principal_cargo"], "")
        self.assertEqual(values["contacto_alternativo_cargo"], "Otro cargo")

    def test_same_person_can_supply_missing_fields_from_prior_year(self):
        values, _ = prepare_company(example_row(N="María  Pérez", AE="MARIA PEREZ", AF="Gerente"))
        self.assertEqual(values["contacto_principal_nombre"], "María Pérez")
        self.assertEqual(values["contacto_principal_cargo"], "Gerente")
        self.assertEqual(values["contacto_alternativo_nombre"], "")

    def test_company_and_sunafil_channels_are_used_and_preserved(self):
        values, source = prepare_company(example_row(AC=123456, AB="044", BD="987654321, 976543210",
                                                    BC="first@example.com;second@example.com"))
        self.assertEqual(values["contacto_principal_telefono"], "44123456")
        self.assertEqual(values["contacto_principal_correo"], "first@example.com")
        self.assertEqual(len(source["telefonos"]), 3)
        self.assertEqual(len(source["correos"]), 2)

    def test_lists_are_split_and_whatsapp_never_concatenates_numbers(self):
        self.assertEqual(phones("987654321 - 976543210"), ["987654321", "976543210"])
        self.assertEqual(phones("987654321\n976543210"), ["987654321", "976543210"])
        self.assertEqual(whatsapp_url("987654321 - 976543210"), "https://wa.me/51987654321")
        self.assertEqual(whatsapp_url("44123456"), "")
        self.assertEqual(emails("first@example.com/second@example.com"), ["first@example.com", "second@example.com"])
        self.assertEqual(phones("0"), [])
        self.assertEqual(phones("-"), [])


class CargaImportTests(TestCase):
    def test_import_is_idempotent_and_keeps_original_zero_separate_from_contact_coverage(self):
        row = example_row(AE="Persona de 2025", AH="987654321")
        company, result = Command().import_row(row, {"fila": 4})
        self.assertEqual(result, "nueva")
        self.assertEqual(company.estado_encuesta, EstadoEncuesta.POR_CONTACTAR)
        self.assertEqual(company.situacion_empresa, "")
        self.assertEqual(company.datos_carga["columnas"]["W"], "0")
        _, result = Command().import_row(row, {"fila": 4})
        self.assertEqual(result, "sin_cambios")
        self.assertEqual(company.history.count(), 1)

    def test_local_edits_and_progress_are_preserved(self):
        row = example_row(N="Persona original", R="original@example.com")
        company, _ = Command().import_row(row, {})
        company.contacto_principal_nombre = "Persona editada"
        company.contacto_principal_correo = "local@example.com"
        company.estado_encuesta = EstadoEncuesta.EN_SEGUIMIENTO
        company.save()
        row.update(N="Nueva persona en fuente", R="new@example.com", O="Cargo nuevo")
        company, _ = Command().import_row(row, {})
        self.assertEqual(company.contacto_principal_nombre, "Persona editada")
        self.assertEqual(company.contacto_principal_correo, "local@example.com")
        self.assertEqual(company.contacto_principal_cargo, "")
        self.assertEqual(company.estado_encuesta, EstadoEncuesta.EN_SEGUIMIENTO)

    def fake_workbook(self, rows):
        header = [""] * 56
        for column, label in LAYOUT.items():
            header[column_index_from_string(column) - 1] = label
        data = []
        for row in rows:
            cells = [None] * 56
            for column, value in row.items():
                cells[column_index_from_string(column) - 1] = value
            data.append(tuple(cells))
        sheet = Mock(title="Sheet1")
        sheet.iter_rows.side_effect = lambda **kwargs: iter([tuple(header)] if kwargs.get("min_row") == 3 else data)
        return Mock(active=sheet)

    def test_dry_run_does_not_persist(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            path.touch()
            with patch("bitacora.management.commands.importar_carga.load_workbook", return_value=self.fake_workbook([example_row()])):
                call_command("importar_carga", path, dry_run=True, stdout=StringIO())
        self.assertFalse(Empresa.objects.exists())

    def test_duplicate_ruc_rolls_back_entire_import(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            path.touch()
            with patch("bitacora.management.commands.importar_carga.load_workbook", return_value=self.fake_workbook([example_row(), example_row()])), self.assertRaises(CommandError):
                call_command("importar_carga", path, stdout=StringIO())
        self.assertFalse(Empresa.objects.exists())

    def test_dashboard_distinguishes_contact_data_from_situation_zero(self):
        Command().import_row(example_row(Q="987654321"), {})
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["con_contacto_count"], 1)
        self.assertEqual(response.context["codigo_cero_count"], 1)
        self.assertEqual(response.context["situaciones"][0]["label"], "Situación sin clasificar")

    def test_legacy_import_is_completed_without_losing_fixed_phone(self):
        Empresa.objects.create(ruc="20000000001", razon_social="Empresa de prueba", region="Lima",
                               ciiu="7492", contacto_principal_telefono="1234567")
        company, _ = Command().import_row(example_row(P="1234567", AE="Persona anterior", AG="old@example.com"), {})
        self.assertEqual(company.contacto_principal_nombre, "Persona anterior")
        self.assertEqual(company.contacto_principal_telefono, "1234567")
        self.assertEqual(company.contacto_principal_correo, "old@example.com")
        self.assertEqual(company.estado_encuesta, EstadoEncuesta.POR_CONTACTAR)
        response = self.client.get(f"/empresas/{company.pk}/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Fuentes de contacto de la carga")
        self.assertContains(response, "EDO 2025")

    def test_alternate_contact_counts_once_per_company(self):
        company, _ = Command().import_row(example_row(), {})
        company.contacto_alternativo_estado = EstadoContacto.CONTACTADO
        company.contacto_alternativo_correo = "alternate@example.com"
        company.save()
        response = self.client.get("/")
        self.assertEqual(response.context["contactadas_count"], 1)
        self.assertEqual(response.context["empresas_sin_correo_count"], 0)
        self.assertEqual(sum(item["total"] for item in response.context["reporte"]), 1)
        company.contacto_principal_estado = EstadoContacto.CONTACTADO
        company.save()
        self.assertEqual(self.client.get("/").context["contactadas_count"], 1)


class DemoEntregableTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        Empresa.objects.bulk_create([
            Empresa(ruc=f"20{index:09d}", razon_social=f"Empresa de prueba {index}",
                    region=f"Región {index % 5}", ciiu="7492",
                    contacto_principal_nombre=f"Informante de prueba {index}",
                    contacto_principal_telefono=f"987{index:06d}",
                    contacto_principal_correo=f"contacto{index}@example.com",
                    estado_encuesta=EstadoEncuesta.POR_CONTACTAR)
            for index in range(1, 182)
        ])

    @override_settings(DEMO_MODE=False)
    def test_real_mode_rejects_demo_seed(self):
        with self.assertRaises(CommandError):
            call_command("preparar_demo", stdout=StringIO())
        self.assertEqual(Actividad.objects.count(), 0)

    @override_settings(DEMO_MODE=True)
    def test_37_contacts_would_exceed_20_percent_and_are_rejected(self):
        with self.assertRaises(CommandError):
            call_command("preparar_demo", contactadas=37, stdout=StringIO())
        self.assertEqual(Actividad.objects.count(), 0)
        self.assertEqual(Empresa.objects.filter(contacto_principal_estado=EstadoContacto.CONTACTADO).count(), 0)

    @override_settings(DEMO_MODE=True)
    def test_scenario_counts_distinct_companies_and_three_days(self):
        call_command("preparar_demo", stdout=StringIO())
        self.assertEqual(Empresa.objects.filter(contacto_principal_estado=EstadoContacto.CONTACTADO).count(), 36)
        self.assertEqual(Empresa.objects.filter(contacto_principal_estado=EstadoContacto.NO_CONTESTA).count(), 9)
        self.assertEqual(Actividad.objects.values("empresa_id").distinct().count(), 45)
        self.assertEqual(Actividad.objects.count(), 57)
        self.assertEqual(len(set(Actividad.objects.values_list("fecha", flat=True))), 3)
        self.assertEqual(Critica.objects.count(), 3)
        response = self.client.get("/")
        self.assertContains(response, "Gestiones y estados simulados")
        self.assertEqual(response.context["asistencia_count"], 12)
        self.assertEqual(sum(item["total"] for item in response.context["reporte"]), 181)
        self.assertLessEqual(response.context["contactadas_porcentaje"], 20)
        export = self.client.get("/actividades/?export=csv")
        self.assertIn("DEMO-", export["Content-Disposition"])
        self.assertIn("Tipo de datos", export.content.decode("utf-8"))
        call_command("preparar_demo", stdout=StringIO())
        self.assertEqual(Actividad.objects.count(), 57)
