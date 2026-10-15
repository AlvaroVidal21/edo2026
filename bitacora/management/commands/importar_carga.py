import re
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from bitacora.carga import clean_text, prepare_company
from bitacora.models import (
    Empresa,
    EstadoContacto,
    EstadoCritica,
    EstadoEncuesta,
    SituacionEmpresa,
)

LAYOUT = {
    "E": "Región", "F": "RUC", "G": "RAZÓN SOCIAL", "H": "codciiu_2025",
    "J": "Correo Sistema", "K": "CORREO PRIORITARIO", "L": "CORREOS UNICOS CONCATENADOS",
    "M": "CORREO PRIORITARIO 2", "N": "Nombres y Apellidos", "O": "Cargo",
    "P": "Teléfono fijo", "Q": "Teléfono Celular", "R": "Correo Eléctronico",
    "S": "Otros Correo Eléctronico", "T": "Correo de la empresa", "U": "Teléfono",
    "W": "CODIGO SITUACIÓN", "Z": "CIIU", "AB": "CÓDIGO DEL TELÉFONO",
    "AC": "TELÉFONO DE LA EMPRESA", "AD": "CORREO DE LA EMPRESA",
    "AE": "Nombre y apellido informante 1", "AF": "Cargo informante 1",
    "AG": "Correo electrónico informante 1", "AH": "Celular informante 1",
    "AI": "Código ciudad informante 1", "AJ": "Teléfono informante 1",
    "AK": "Nombre y apellido informante 2", "AL": "Cargo informante 2",
    "AM": "Correo electrónico informante 2", "AN": "Celular informante 2",
    "AO": "Código ciudad informante 2", "AP": "Teléfono informante 2",
    "AQ": "Nombre y apellido informante 1", "AR": "Cargo informante 1",
    "AS": "Correo electrónico informante 1", "AT": "Celular informante 1",
    "AU": "Código ciudad informante 1", "AV": "Teléfono informante 1",
    "AW": "Nombre y apellido informante 2", "AX": "Cargo informante 2",
    "AY": "Correo electrónico informante 2", "AZ": "Celular informante 2",
    "BA": "Código ciudad informante 2", "BB": "Teléfono informante 2",
    "BC": "Correos electrónicos", "BD": "Teléfonos",
}


def legacy_values(row):
    """Identifica los valores de la primera carga para conservar ediciones posteriores."""
    def original_text(value):
        if value is None or value == 0:
            return ""
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    email = ""
    for column in ["R", "K", "J", "M", "S", "T"]:
        candidate = original_text(row.get(column))
        try:
            validate_email(candidate)
        except ValidationError:
            continue
        email = candidate
        break
    return {
        "razon_social": original_text(row.get("G")),
        "region": original_text(row.get("E")) or "Sin dato",
        "ciiu": original_text(row.get("Z")) or original_text(row.get("H")) or "Sin dato",
        "contacto_principal_nombre": original_text(row.get("N")),
        "contacto_principal_cargo": original_text(row.get("O")),
        "contacto_principal_telefono": original_text(row.get("Q")) or original_text(row.get("P")),
        "contacto_principal_correo": email,
        **{f"contacto_alternativo_{field}": "" for field in ["nombre", "cargo", "telefono", "correo"]},
    }


class Command(BaseCommand):
    help = "Cruza todas las fuentes del Excel EDO conservando originales y procedencia."

    def add_arguments(self, parser):
        parser.add_argument("archivo", type=Path)
        parser.add_argument("--dry-run", action="store_true", help="Revisa el resultado sin guardar.")

    def import_row(self, row, file_info):
        ruc = clean_text(row.get("F"))
        values, metadata = prepare_company(row)
        metadata.update(file_info)
        company = Empresa.objects.filter(ruc=ruc).first()
        created = company is None
        if created:
            company = Empresa(ruc=ruc, **values)
        else:
            previous = company.datos_carga.get("valores_importados")
            baseline = previous if previous is not None else legacy_values(row)
            protected_prefixes = set()
            for prefix in ["contacto_principal", "contacto_alternativo"]:
                name_field = f"{prefix}_nombre"
                if clean_text(getattr(company, name_field)) != clean_text(baseline.get(name_field)):
                    protected_prefixes.add(prefix)
            for field, value in values.items():
                current = getattr(company, field)
                protected = any(field.startswith(prefix + "_") for prefix in protected_prefixes)
                if not protected and clean_text(current) == clean_text(baseline.get(field)):
                    setattr(company, field, value)
                elif current != value:
                    metadata["origenes"][field] = "Edición local conservada"

        has_channel = any(
            getattr(company, field) for field in ["contacto_principal_correo", "contacto_principal_telefono",
                                                "contacto_alternativo_correo", "contacto_alternativo_telefono"]
        )
        untouched_legacy = (
            not created and not company.datos_carga
            and company.estado_encuesta == EstadoEncuesta.BUSQUEDA
            and company.contacto_principal_estado == EstadoContacto.PENDIENTE
            and company.estado_critica == EstadoCritica.SIN_INICIAR
            and not company.actividades.exists() and not company.criticas.exists()
        )
        if (created or untouched_legacy) and has_channel:
            company.estado_encuesta = EstadoEncuesta.POR_CONTACTAR
        if company.contacto_alternativo_nombre and not company.contacto_alternativo_estado:
            company.contacto_alternativo_estado = EstadoContacto.PENDIENTE

        situation = clean_text(row.get("W"))
        if situation in SituacionEmpresa.values and (company.situacion_empresa != situation or not company.situacion_actualizada_en):
            company.situacion_empresa = situation
            company.situacion_actualizada_en = timezone.localdate()
        # El código 0 se conserva en el original; no se inventa una equivalencia.
        company.datos_carga = metadata
        company._change_reason = "Cruce de carga: actualizados, base, Planilla 2025, EDO 2025/2024 y SUNAFIL."
        if created:
            company.full_clean()
            company.save()
            return company, "nueva"
        old = Empresa.objects.get(pk=company.pk)
        changed = any(getattr(old, field.attname) != getattr(company, field.attname)
                      for field in Empresa._meta.concrete_fields if not field.primary_key)
        if changed:
            company.full_clean()
            company.save()
        return company, "actualizada" if changed else "sin_cambios"

    def handle(self, *args, **options):
        path = options["archivo"]
        if not path.is_file():
            raise CommandError("No se encuentra el archivo de carga.")
        workbook = load_workbook(path, read_only=True, data_only=True)
        counts = {"nueva": 0, "actualizada": 0, "sin_cambios": 0, "con_correo": 0, "con_telefono": 0, "con_nombre": 0}
        try:
            sheet = workbook.active
            header = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
            headers = {get_column_letter(i): clean_text(value) for i, value in enumerate(header, 1)}
            for column, expected in LAYOUT.items():
                if headers.get(column) != clean_text(expected):
                    raise CommandError(f"Estructura inesperada en {column}3. No se modificó el directorio.")
            seen = set()
            with transaction.atomic():
                for number, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
                    if not any(value is not None for value in values):
                        continue
                    row = {get_column_letter(i): value for i, value in enumerate(values, 1)}
                    ruc = clean_text(row.get("F"))
                    if not re.fullmatch(r"\d{11}", ruc) or not clean_text(row.get("G")):
                        raise CommandError(f"Identificación inválida en la fila {number}. No se guardó la carga.")
                    if ruc in seen:
                        raise CommandError(f"RUC repetido en la fila {number}. No se guardó la carga.")
                    seen.add(ruc)
                    info = {"archivo": path.name, "hoja": sheet.title, "fila": number, "encabezados": headers}
                    company, outcome = self.import_row(row, info)
                    counts[outcome] += 1
                    for key, field in [("con_correo", "contacto_principal_correo"), ("con_telefono", "contacto_principal_telefono"), ("con_nombre", "contacto_principal_nombre")]:
                        counts[key] += bool(getattr(company, field))
                if options["dry_run"]:
                    transaction.set_rollback(True)
        finally:
            workbook.close()
        prefix = "Simulación sin guardar" if options["dry_run"] else "Carga guardada"
        self.stdout.write(self.style.SUCCESS(f"{prefix}: {counts}."))
